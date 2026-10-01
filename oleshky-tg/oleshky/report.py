"""Етап 4 — CSV/XLSX, офлайн HTML-звіт, markdown для NotebookLM."""
from __future__ import annotations

import csv
import logging
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

from jinja2 import Environment, FileSystemLoader, select_autoescape

from . import config, db
from .classify import final_relevance
from .filters import mask_pii

log = logging.getLogger("oleshky.report")
KYIV = ZoneInfo(config.KYIV_TZ_NAME)

POST_COLUMNS = ["date", "channel", "side", "link", "type", "blockade_related", "summary_uk", "text",
                "matched_terms", "views", "forwards", "repost_count", "relevance"]
CHANNEL_COLUMNS = ["username", "title", "type", "participants", "side", "relevant", "local",
                   "posts_matched", "last_post_date", "found_via"]
TYPE_LABELS = {
    "appeal": "Звернення", "testimony": "Свідчення", "humanitarian": "Гуманітарна ситуація",
    "evacuation": "Евакуація / виїзд", "shelling": "Обстріли", "occupation_admin": "Окупаційна адміністрація",
    "propaganda": "Пропаганда", "news": "Новини", "other": "Інше", "": "Не класифіковано",
}
SIDE_LABELS = {"ua": "UA", "occupation": "окупація", "neutral": "нейтральні", "unknown": "невідомо"}
NOTEBOOK_CHUNK = 300_000   # символів на файл NotebookLM


def kyiv(iso: str | None, fmt: str = "%Y-%m-%d %H:%M") -> str:
    if not iso:
        return ""
    return datetime.fromisoformat(iso).astimezone(KYIV).strftime(fmt)


def load_posts(conn, labels, *, since_last: bool, since: str | None) -> list[dict]:
    chans = {r["id"]: db.effective_channel(r, labels) for r in conn.execute("SELECT * FROM channels")}
    reported = {r[0] for r in conn.execute("SELECT post_key FROM reported_posts")} if since_last else set()
    rows = conn.execute("""
        SELECT p.*, cl.relevant AS llm_relevant, cl.type AS llm_type, cl.blockade_related,
               cl.side AS llm_side, cl.summary_uk, cl.evidence, cl.evidence_ok
        FROM posts p LEFT JOIN classifications cl ON cl.text_hash = p.text_hash
        WHERE p.relevance_rule IN ('yes','unsure') ORDER BY p.date DESC""").fetchall()

    group_channels: dict[str, set] = defaultdict(set)
    for r in conn.execute("SELECT repost_group, channel_id FROM posts WHERE repost_group IS NOT NULL"):
        group_channels[r[0]].add(r[1])

    out = []
    for r in rows:
        key = db.post_key(r["channel_id"], r["msg_id"])
        if key in reported:
            continue
        if since and r["date"] < since:
            continue
        ch = chans.get(r["parent_channel_id"] if r["is_comment"] else r["channel_id"], {})
        d = dict(r)
        d["key"] = key
        d["channel"] = f"@{ch['username']}" if ch.get("username") else str(r["channel_id"])
        d["channel_username"] = ch.get("username")
        d["channel_side"] = ch.get("side", "unknown")
        d["side"] = r["llm_side"] if r["llm_side"] not in (None, "unknown") else ch.get("side", "unknown")
        d["type"] = r["llm_type"] or ""
        d["relevance"] = final_relevance(r["relevance_rule"], r["llm_relevant"])
        d["repost_count"] = max(0, len(group_channels.get(r["repost_group"], ())) - 1)
        d["date_kyiv"] = kyiv(r["date"])
        d["day"] = kyiv(r["date"], "%Y-%m-%d")
        if r["llm_relevant"] is None:
            d["unsure_reason"] = "лише правило (LLM ще не запускали)"
        elif r["llm_relevant"] == "unsure" and r["evidence_ok"] == 0:
            d["unsure_reason"] = "цитата LLM не знайдена в тексті"
        else:
            d["unsure_reason"] = "LLM не впевнена"
        out.append(d)
    return out


def _mask(text, public, enabled):
    return mask_pii(text or "", public) if enabled else (text or "")


def write_csv(path: Path, rows: list[dict], columns: list[str]):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def write_xlsx(path: Path, rows: list[dict], columns: list[str]):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "posts"
    ws.append(columns)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([r.get(c) for c in columns])
    widths = {"date": 17, "channel": 22, "side": 11, "link": 38, "type": 16, "blockade_related": 9,
              "summary_uk": 60, "text": 90, "matched_terms": 18, "views": 9, "forwards": 9,
              "repost_count": 9, "relevance": 9}
    for i, c in enumerate(columns, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = widths.get(c, 14)
    link_col = columns.index("link") + 1
    for row in ws.iter_rows(min_row=2):
        cell = row[link_col - 1]
        if cell.value:
            cell.hyperlink = cell.value
            cell.style = "Hyperlink"
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=columns[c.column - 1] in ("summary_uk", "text"))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def daily_bars(posts: list[dict]) -> dict:
    counts = Counter(p["day"] for p in posts)
    if not counts:
        return {"bars": [], "max": 0, "width": 0}
    days = sorted(counts)
    start, end = date.fromisoformat(days[0]), date.fromisoformat(days[-1])
    n = (end - start).days + 1
    all_days = [(start + timedelta(days=i)).isoformat() for i in range(n)]
    mx = max(counts.values())
    bw = 8
    bars = []
    for i, d in enumerate(all_days):
        v = counts.get(d, 0)
        h = round(100 * v / mx) if mx else 0
        bars.append({"x": i * (bw + 2), "y": 120 - h, "h": h, "day": d, "v": v, "w": bw})
    return {"bars": bars, "max": mx, "width": n * (bw + 2), "first": all_days[0], "last": all_days[-1]}


def write_notebooklm(out_dir: Path, posts: list[dict], public, mask: bool) -> list[Path]:
    """Тільки relevant=true, хронологічно, кластер репостів — один раз. Коментарі не включаємо."""
    nb = out_dir / "notebooklm"
    nb.mkdir(parents=True, exist_ok=True)
    chosen: dict[str, dict] = {}
    for p in sorted(posts, key=lambda p: p["date"]):
        if p["relevance"] != "true" or p["is_comment"]:
            continue
        g = p["repost_group"] or p["key"]
        chosen.setdefault(g, p)
    by_month: dict[str, list[dict]] = defaultdict(list)
    for p in sorted(chosen.values(), key=lambda p: p["date"]):
        by_month[kyiv(p["date"], "%Y-%m")].append(p)

    files = []
    for month, items in sorted(by_month.items()):
        header = (f"# Олешки — Telegram, {month}\n\n"
                  "Добірка постів із публічних Telegram-каналів, де згадуються Олешки (Херсонщина). "
                  "Дати — за Києвом. Телефони, email і @username приватних акаунтів замасковано.\n\n")
        chunks, cur = [], header
        for p in items:
            title = f"## {kyiv(p['date'])} · {p['channel']} ({p['side']}) · {p['type'] or 'не класифіковано'}"
            lines = [title, f"Посилання: {p['link']}"]
            if p["repost_count"]:
                lines.append(f"Репостнули: {p['repost_count']} каналів")
            if p.get("summary_uk"):
                lines.append(f"Коротко: {_mask(p['summary_uk'], public, mask)}")
            lines.append("")
            lines.append(_mask(p["text"], public, mask).strip())
            entry = "\n".join(lines) + "\n\n"
            if len(cur) + len(entry) > NOTEBOOK_CHUNK and cur != header:
                chunks.append(cur)
                cur = header
            cur += entry
        chunks.append(cur)
        for i, text in enumerate(chunks, 1):
            name = f"oleshky_tg_{month}.md" if len(chunks) == 1 else f"oleshky_tg_{month}_part{i}.md"
            path = nb / name
            path.write_text(text, encoding="utf-8")
            files.append(path)
    return files


def run_report(conn, labels, *, since_last: bool = False, mask: bool = True, since: str | None = None,
               out_dir: Path | None = None) -> Path:
    now_kyiv = datetime.now(KYIV)
    out_dir = out_dir or config.EXPORTS_DIR / now_kyiv.strftime("%Y-%m-%d")
    out_dir.mkdir(parents=True, exist_ok=True)
    public = db.public_usernames(conn)
    posts = load_posts(conn, labels, since_last=since_last, since=since)

    # ---- posts.csv / xlsx (LLM-false не виключаємо — колонка relevance)
    rows = []
    for p in posts:
        rows.append({
            "date": p["date_kyiv"], "channel": p["channel"], "side": p["side"], "link": p["link"],
            "type": p["type"], "blockade_related": "" if p["blockade_related"] is None else bool(p["blockade_related"]),
            "summary_uk": _mask(p["summary_uk"], public, mask), "text": _mask(p["text"], public, mask),
            "matched_terms": p["matched_terms"], "views": p["views"], "forwards": p["forwards"],
            "repost_count": p["repost_count"], "relevance": p["relevance"],
        })
    write_csv(out_dir / "posts.csv", rows, POST_COLUMNS)
    write_xlsx(out_dir / "posts.xlsx", rows, POST_COLUMNS)

    # ---- channels.csv
    stats = {r[0]: (r[1], r[2]) for r in conn.execute(
        "SELECT COALESCE(parent_channel_id, channel_id), COUNT(*), MAX(date) FROM posts "
        "WHERE relevance_rule='yes' GROUP BY 1")}
    reported_ch = {r[0] for r in conn.execute("SELECT channel_id FROM reported_channels")}
    channels, new_channels = [], []
    for row in db.get_channels(conn):
        ch = db.effective_channel(row, labels)
        n, last = stats.get(ch["id"], (0, None))
        item = {
            "username": ch["username"], "title": ch["title"], "type": ch["type"],
            "participants": ch["participants"], "side": ch["side"],
            "relevant": "" if ch["relevant"] is None else ch["relevant"], "local": ch["local"],
            "posts_matched": n, "last_post_date": kyiv(last), "found_via": ch["found_via"],
        }
        channels.append(item)
        if ch["id"] not in reported_ch and ch["relevant"] is not False:
            new_channels.append(item)
    channels.sort(key=lambda c: (-c["posts_matched"], c["username"].lower()))
    write_csv(out_dir / "channels.csv", channels, CHANNEL_COLUMNS)

    # ---- HTML
    relevant = [p for p in posts if p["relevance"] == "true"]
    for p in posts:
        p["text_masked"] = _mask(p["text"], public, mask)
        p["summary_masked"] = _mask(p["summary_uk"], public, mask)
    seen_groups, blockade = set(), []
    for p in sorted((p for p in relevant if p["blockade_related"]), key=lambda p: -(p["views"] or 0)):
        g = p["repost_group"] or p["key"]
        if g not in seen_groups:
            seen_groups.add(g)
            blockade.append(p)
    ctx = {
        "generated": now_kyiv.strftime("%Y-%m-%d %H:%M"),
        "since_last": since_last,
        "masked": mask,
        "period": (min(p["date_kyiv"] for p in relevant)[:10], max(p["date_kyiv"] for p in relevant)[:10]) if relevant else None,
        "n_relevant": len(relevant),
        "n_unique": len({p["repost_group"] or p["key"] for p in relevant}),
        "n_unsure": sum(1 for p in posts if p["relevance"] == "unsure"),
        "n_channels": len({p["channel"] for p in relevant}),
        "daily": daily_bars(relevant),
        "by_type": [(TYPE_LABELS.get(k, k), v) for k, v in Counter(p["type"] for p in relevant).most_common()],
        "by_side": [(SIDE_LABELS.get(k, k), v) for k, v in Counter(p["side"] for p in relevant).most_common()],
        "appeals": sorted([p for p in relevant if p["type"] in ("appeal", "testimony")], key=lambda p: p["date"], reverse=True),
        "blockade": blockade[:30],
        "new_channels": new_channels,
        "manual": [p for p in posts if p["relevance"] == "unsure"][:300],
        "type_labels": TYPE_LABELS,
        "side_labels": SIDE_LABELS,
    }
    env = Environment(loader=FileSystemLoader(Path(__file__).parent / "templates"),
                      autoescape=select_autoescape(["html"]))
    (out_dir / "report.html").write_text(env.get_template("report.html").render(**ctx), encoding="utf-8")

    nb_files = write_notebooklm(out_dir, posts, public, mask)

    # ---- позначаємо як «уже у звіті»
    today = now_kyiv.strftime("%Y-%m-%d")
    conn.executemany("INSERT OR IGNORE INTO reported_posts (post_key, report_date) VALUES (?,?)",
                     [(p["key"], today) for p in posts])
    conn.executemany("INSERT OR IGNORE INTO reported_channels (channel_id, report_date) VALUES (?,?)",
                     [(r["id"], today) for r in db.get_channels(conn)])
    conn.commit()
    log.info("Звіт: %s (постів %d, релевантних %d, перевірити %d, NotebookLM файлів %d)",
             out_dir, len(posts), len(relevant), ctx["n_unsure"], len(nb_files))
    return out_dir
